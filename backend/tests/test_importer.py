import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from app.services import importer
from app.services.importer import ImportFileError, read_table, suggest_mapping, validate_rows


def make_xlsx(rows, title_row=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    if title_row:
        ws.append([title_row])
        ws.append([])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE = [
    ["Nom", "Prénom", "Téléphone", "Montant", "Date échéance", "Référence"],
    ["Alaoui", "Salma", "0612345678", 1250.5, datetime(2026, 11, 5), "FAC-1"],
    ["Bennani", "Youssef", 661234567, 800, datetime(2026, 11, 6), "FAC-2"],
    ["Tazi", "Omar", "06 12 34 56 78", 300, datetime(2026, 11, 7), "FAC-3"],   # doublon de la ligne 2
    ["Fixe", "Test", "0522123456", 10, None, "FAC-4"],                         # fixe
    ["Vide", "X", None, 10, None, "FAC-5"],                                    # vide
]


def test_read_xlsx_detects_header_after_title():
    table = read_table(make_xlsx(SAMPLE, title_row="Export clients octobre"), "clients.xlsx")
    assert table.columns == SAMPLE[0]
    assert table.header_row == 3
    assert len(table.rows) == 5
    assert table.rows[1]["Téléphone"] == 661234567


def test_suggest_mapping_by_header():
    table = read_table(make_xlsx(SAMPLE), "c.xlsx")
    mapping = suggest_mapping(table)
    assert mapping["Téléphone"] == importer.PHONE_FIELD
    assert mapping["Date échéance"] == "date_echeance"
    assert mapping["Prénom"] == "prenom"


def test_suggest_mapping_by_content():
    rows = [["Client", "Contact"], ["A", "0612345678"], ["B", "0700112233"]]
    mapping = suggest_mapping(read_table(make_xlsx(rows), "c.xlsx"))
    assert mapping["Contact"] == importer.PHONE_FIELD


def test_validate_rows_statuses_and_render():
    table = read_table(make_xlsx(SAMPLE), "c.xlsx")
    report = validate_rows(table, suggest_mapping(table),
                           "Bonjour {prenom}, {montant} MAD avant le {date_echeance} ({reference})")
    s = report.summary()
    assert s == {"total": 5, "valid": 2, "invalid": 2, "duplicate": 1, "blacklisted": 0}
    first = report.rows[0]
    assert first.phone_e164 == "+212612345678"
    assert first.rendered == "Bonjour Salma, 1 250,50 MAD avant le 05/11/2026 (FAC-1)"
    assert report.rows[1].phone_e164 == "+212661234567"  # zéro perdu par Excel
    assert report.rows[2].status == importer.ROW_DUPLICATE and "ligne 2" in report.rows[2].error
    assert report.rows[3].status == importer.ROW_INVALID and "fixe" in report.rows[3].error.lower()


def test_blacklist_and_no_dedupe():
    table = read_table(make_xlsx(SAMPLE), "c.xlsx")
    report = validate_rows(table, suggest_mapping(table), "Hello", blacklist={"+212661234567"}, dedupe=False)
    assert report.rows[1].status == importer.ROW_BLACKLISTED
    assert report.rows[2].status == importer.ROW_VALID


def test_message_column_overrides_body():
    rows = [["tel", "message", "nom"], ["0612345678", "Spécial {nom}", "A"], ["0700112233", None, "B"]]
    table = read_table(make_xlsx(rows), "c.xlsx")
    mapping = suggest_mapping(table)
    assert mapping["message"] == importer.MESSAGE_FIELD
    report = validate_rows(table, mapping, "Commun {nom}")
    assert [r.rendered for r in report.rows] == ["Spécial A", "Commun B"]
    # Sans texte commun : la ligne sans message est invalide
    report = validate_rows(table, mapping, None)
    assert report.rows[1].status == importer.ROW_INVALID


def test_csv_semicolon_and_arabic_cp1256():
    content = "الاسم;الهاتف\nفاطمة;0612345678\nعمر;0700112233\n".encode("cp1256")
    table = read_table(content, "liste.csv")
    assert table.columns == ["الاسم", "الهاتف"]
    mapping = suggest_mapping(table)
    assert mapping["الهاتف"] == importer.PHONE_FIELD
    report = validate_rows(table, mapping, "مرحبا {prenom}")  # « الاسم » est un alias de prenom
    assert report.rows[0].rendered == "مرحبا فاطمة"


def test_csv_comma_utf8_bom():
    content = "﻿phone,name\n+212612345678,Ali\n".encode("utf-8")
    table = read_table(content, "x.csv")
    assert table.columns == ["phone", "name"]


def test_errors():
    with pytest.raises(ImportFileError):
        read_table(b"", "x.xlsx")
    with pytest.raises(ImportFileError):
        read_table(b"garbage", "x.xlsx")
    with pytest.raises(ImportFileError):
        read_table(b"a", "x.pdf")
    with pytest.raises(ImportFileError, match="too_many_rows"):
        read_table(make_xlsx([["tel"]] + [["0612345678"]] * 30), "x.xlsx", max_rows=10)


def test_error_report_xlsx():
    table = read_table(make_xlsx(SAMPLE), "c.xlsx")
    report = validate_rows(table, suggest_mapping(table), "Hi")
    data = importer.build_error_report(report, table.columns)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    assert ws.max_row == 1 + 3  # en-tête + 3 lignes rejetées
    assert ws.cell(2, 1).value == 4  # numéro de ligne dans le fichier
    assert wb.worksheets[1].cell(1, 2).value == 5
