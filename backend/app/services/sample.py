"""Fichier Excel d'exemple pour l'import (aussi livré dans samples/import_exemple.xlsx)."""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["Nom", "Prénom", "Téléphone", "Montant", "Date échéance", "Référence", "Ville", "Langue", "Message"]


def sample_rows(today: date | None = None) -> list[list]:
    d = today or date.today()
    return [
        ["Alaoui", "Salma", "0612345678", 1250.50, d + timedelta(days=5), "FAC-2026-001", "Casablanca", "fr", None],
        ["Bennani", "Youssef", "06 61 23 45 67", 800, d + timedelta(days=7), "FAC-2026-002", "Rabat", "fr", None],
        ["El Idrissi", "Fatima", "+212700112233", 450, d + timedelta(days=3), "FAC-2026-003", "Fès", "ar",
         "مرحبا {prenom}، نذكركم بأن مبلغ {montant} درهم يحل أجله في {date_echeance}."],
        ["Tazi", "Omar", "0655443322", 2300, d - timedelta(days=4), "FAC-2026-004", "Marrakech", "fr",
         "Bonjour {prenom}, votre facture {reference} de {montant} MAD est échue depuis le {date_echeance}. Merci de régulariser."],
        ["Chraibi", "Nadia", "212661112233", 990, d + timedelta(days=10), "FAC-2026-005", "Tanger", "fr", None],
        ["Doublon", "Test", "06.12.34.56.78", 100, d, "FAC-2026-006", "Casablanca", "fr", None],
        ["Fixe", "Test", "0522123456", 100, d, "FAC-2026-007", "Casablanca", "fr", None],
        ["Invalide", "Test", "12345", 100, d, "FAC-2026-008", "Agadir", "fr", None],
    ]


def build_sample_xlsx(today: date | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"
    ws.append(HEADERS)
    for row in sample_rows(today):
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "@"          # téléphone en texte (garde le 0 initial)
        row[3].number_format = "#,##0.00"
        row[4].number_format = "DD/MM/YYYY"
        row[8].alignment = Alignment(wrap_text=True)
    widths = [14, 12, 17, 12, 15, 15, 13, 9, 70]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    ws["C1"].comment = Comment("Formats acceptés : 06…, 07…, +212…, 00212…, avec ou sans espaces/points.", "SMS Gateway")
    ws["I1"].comment = Comment("Facultatif : message propre à la ligne. Vide = texte commun saisi dans l'assistant.",
                               "SMS Gateway")
    dv = DataValidation(type="list", formula1='"fr,ar"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H2:H1000")

    help_ws = wb.create_sheet("Mode d'emploi")
    lines = [
        "IMPORT DE CONTACTS ET ENVOI DE SMS PERSONNALISÉS",
        "",
        "1. Une ligne = un destinataire. La première ligne contient les noms de colonnes.",
        "2. Colonne téléphone obligatoire : 06/07…, +212…, 00212… (les fixes 05… sont refusés).",
        "3. Chaque colonne devient une variable : « Date échéance » -> {date_echeance}, « Ville » -> {ville}.",
        "4. Exemple de texte commun : Bonjour {prenom}, votre échéance {reference} de {montant} MAD est due le {date_echeance}.",
        "5. Colonne « Message » (facultative) : texte spécifique à la ligne, prioritaire sur le texte commun.",
        "6. Les doublons, numéros invalides et désinscrits (STOP) sont détectés avant l'envoi ;",
        "   un rapport d'erreurs Excel est téléchargeable.",
        "7. Colonnes Montant + Date échéance + Référence : créent des échéances utilisées par les rappels automatiques.",
        "",
        "Les lignes « Doublon », « Fixe » et « Invalide » de l'onglet Clients illustrent les contrôles.",
    ]
    for line in lines:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True, size=13)
    help_ws.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
