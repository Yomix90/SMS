"""Statistiques, historique filtrable et exports Excel / PDF."""
import io
import uuid as uuidlib
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..extensions import db
from ..models import Campaign, Device, Message
from ..models.message import DIR_IN, DIR_OUT, ST_DELIVERED, ST_FAILED, ST_SENT
from .timeutil import local_day_bounds, local_today, to_local

STATUS_LABELS = {
    "fr": {"scheduled": "Planifié", "queued": "En attente", "paused": "En pause", "dispatched": "Transmis",
           "sent": "Envoyé", "delivered": "Livré", "failed": "Échoué", "canceled": "Annulé",
           "blocked": "Bloqué", "received": "Reçu"},
    "ar": {"scheduled": "مبرمج", "queued": "في الانتظار", "paused": "متوقف", "dispatched": "مرسل للهاتف",
           "sent": "مرسل", "delivered": "مستلم", "failed": "فشل", "canceled": "ملغى", "blocked": "محظور",
           "received": "وارد"},
}


@dataclass
class MessageFilters:
    date_from: date | None = None
    date_to: date | None = None
    status: str | None = None
    direction: str | None = None
    device_id: int | None = None
    campaign_id: int | None = None
    phone: str | None = None
    search: str | None = None
    source: str | None = None
    api_key_id: int | None = None


def filtered_messages(f: MessageFilters):
    q = db.session.query(Message)
    if f.date_from:
        q = q.filter(Message.created_at >= local_day_bounds(f.date_from)[0])
    if f.date_to:
        q = q.filter(Message.created_at < local_day_bounds(f.date_to)[1])
    if f.status:
        statuses = [s.strip() for s in f.status.split(",") if s.strip()]
        q = q.filter(Message.status.in_(statuses))
    if f.direction:
        q = q.filter(Message.direction == f.direction)
    if f.device_id:
        q = q.filter(Message.device_id == f.device_id)
    if f.campaign_id:
        q = q.filter(Message.campaign_id == f.campaign_id)
    if f.phone:
        q = q.filter(Message.phone.contains(f.phone.replace(" ", "")))
    if f.search:
        q = q.filter(Message.body.ilike(f"%{f.search}%"))
    if f.source:
        q = q.filter(Message.source == f.source)
    if f.api_key_id:
        q = q.filter(Message.api_key_id == f.api_key_id)
    return q


def summary(f: MessageFilters) -> dict:
    out = filtered_messages(MessageFilters(**{**f.__dict__, "direction": DIR_OUT}))
    counts = dict(out.with_entities(Message.status, sa.func.count(Message.id)).group_by(Message.status).all())
    segments = out.with_entities(sa.func.coalesce(sa.func.sum(Message.segments), 0)).scalar()
    received = filtered_messages(MessageFilters(**{**f.__dict__, "direction": DIR_IN, "status": None})).count()
    sent_total = counts.get(ST_SENT, 0) + counts.get(ST_DELIVERED, 0)
    finished = sent_total + counts.get(ST_FAILED, 0)
    return {
        "by_status": counts,
        "total_out": sum(counts.values()),
        "sent": sent_total,
        "delivered": counts.get(ST_DELIVERED, 0),
        "failed": counts.get(ST_FAILED, 0),
        "received": received,
        "segments": int(segments or 0),
        "delivery_rate": round(counts.get(ST_DELIVERED, 0) * 100 / sent_total, 1) if sent_total else 0.0,
        "success_rate": round(sent_total * 100 / finished, 1) if finished else 0.0,
    }


def daily_series(days: int = 30, f: MessageFilters | None = None) -> dict:
    """Séries quotidiennes (heure locale) : envoyés, livrés, échoués, reçus."""
    f = f or MessageFilters()
    end = f.date_to or local_today()
    start = f.date_from or (end - timedelta(days=days - 1))
    q = filtered_messages(MessageFilters(**{**f.__dict__, "date_from": start, "date_to": end, "status": None,
                                            "direction": None}))
    rows = q.with_entities(Message.created_at, Message.direction, Message.status).all()
    labels = [(start + timedelta(days=i)) for i in range((end - start).days + 1)]
    index = {d: i for i, d in enumerate(labels)}
    series = {k: [0] * len(labels) for k in ("sent", "delivered", "failed", "received")}
    for created_at, direction, status in rows:
        i = index.get(to_local(created_at).date())
        if i is None:
            continue
        if direction == DIR_IN:
            series["received"][i] += 1
        elif status in (ST_SENT, ST_DELIVERED):
            series["sent"][i] += 1
            if status == ST_DELIVERED:
                series["delivered"][i] += 1
        elif status == ST_FAILED:
            series["failed"][i] += 1
    return {"labels": [d.strftime("%d/%m") for d in labels], **series}


def per_device(f: MessageFilters) -> list[dict]:
    q = (filtered_messages(MessageFilters(**{**f.__dict__, "direction": DIR_OUT}))
         .with_entities(Message.device_id, Message.status, sa.func.count(Message.id))
         .group_by(Message.device_id, Message.status))
    data: dict = {}
    for device_id, status, count in q.all():
        data.setdefault(device_id, {})[status] = count
    names = {d.id: d.name for d in db.session.query(Device).all()}
    return [{"device_id": k, "name": names.get(k, "—"), "sent": v.get(ST_SENT, 0) + v.get(ST_DELIVERED, 0),
             "delivered": v.get(ST_DELIVERED, 0), "failed": v.get(ST_FAILED, 0)} for k, v in data.items()]


# --- Exports -----------------------------------------------------------------------------

EXPORT_COLUMNS = {
    "fr": ["Date", "Sens", "Numéro", "Message", "Statut", "Segments", "Appareil", "SIM", "Campagne",
           "Envoyé le", "Livré le", "Erreur", "Référence"],
    "ar": ["التاريخ", "الاتجاه", "الرقم", "الرسالة", "الحالة", "الأجزاء", "الجهاز", "الشريحة", "الحملة",
           "تاريخ الإرسال", "تاريخ الاستلام", "الخطأ", "المرجع"],
}


def _fmt(dt: datetime | None) -> str:
    return to_local(dt).strftime("%d/%m/%Y %H:%M") if dt else ""


def _rows(q, lang: str, limit: int):
    labels = STATUS_LABELS.get(lang, STATUS_LABELS["fr"])
    devices = {d.id: d.name for d in db.session.query(Device).all()}
    campaigns = dict(db.session.query(Campaign.id, Campaign.name).all())
    direction = {"out": "Sortant" if lang == "fr" else "صادر", "in": "Entrant" if lang == "fr" else "وارد"}
    for m in q.order_by(Message.created_at.desc()).limit(limit).yield_per(500):
        yield [_fmt(m.received_at or m.created_at), direction[m.direction], m.phone, m.body,
               labels.get(m.status, m.status), m.segments, devices.get(m.device_id, ""),
               "" if m.sim_slot is None else f"SIM {m.sim_slot + 1}", campaigns.get(m.campaign_id, ""),
               _fmt(m.sent_at), _fmt(m.delivered_at), m.error_code or "", m.external_ref or ""]


def export_xlsx(f: MessageFilters, lang: str = "fr", limit: int = 100000) -> bytes:
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Messages" if lang == "fr" else "الرسائل"
    if lang == "ar":
        ws.sheet_view.rightToLeft = True
    ws.append(EXPORT_COLUMNS.get(lang, EXPORT_COLUMNS["fr"]))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
    for row in _rows(filtered_messages(f), lang, limit):
        ws.append(row)
    widths = [17, 9, 15, 60, 12, 9, 18, 7, 22, 17, 17, 18, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    s = summary(f)
    ws2 = wb.create_sheet("Résumé" if lang == "fr" else "ملخص")
    for k in ("total_out", "sent", "delivered", "failed", "received", "segments", "delivery_rate"):
        ws2.append([k, s[k]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _shape(text: str) -> str:
    """Mise en forme de l'arabe pour reportlab (liaison des lettres + ordre RTL)."""
    if not any("\u0600" <= c <= "\u06ff" for c in text or ""):
        return text
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except ImportError:
        return text


def _register_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf",
                  str(Path(__file__).resolve().parent.parent / "static" / "fonts" / "DejaVuSans.ttf")]
    for path in candidates:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("AppSans", path))
                return "AppSans"
            except Exception:
                continue
    return "Helvetica"


def export_pdf(f: MessageFilters, lang: str = "fr", limit: int = 2000, title: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _register_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    align = 2 if lang == "ar" else 0
    h = ParagraphStyle("h", fontName=font, fontSize=15, leading=19, alignment=align)
    p = ParagraphStyle("p", fontName=font, fontSize=7.5, leading=9, alignment=align)
    ph = ParagraphStyle("ph", parent=p, textColor=colors.white)
    s = summary(f)
    title = title or ("Rapport des messages" if lang == "fr" else "تقرير الرسائل")
    period = f"{f.date_from or '…'} → {f.date_to or '…'}"
    story = [Paragraph(_shape(title), h), Spacer(1, 3 * mm)]
    kpi_labels = (["Total", "Envoyés", "Livrés", "Échoués", "Reçus", "Taux de livraison"] if lang == "fr"
                  else ["المجموع", "مرسلة", "مستلمة", "فاشلة", "واردة", "نسبة التسليم"])
    kpis = [[_shape(x) for x in kpi_labels],
            [s["total_out"], s["sent"], s["delivered"], s["failed"], s["received"], f"{s['delivery_rate']} %"]]
    t = Table(kpis)
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ECF0F1"))]))
    story += [Paragraph(_shape(period), p), Spacer(1, 2 * mm), t, Spacer(1, 5 * mm)]

    cols = EXPORT_COLUMNS.get(lang, EXPORT_COLUMNS["fr"])
    keep = [0, 2, 3, 4, 6, 9, 10, 11]
    data = [[Paragraph(_shape(cols[i]), ph) for i in keep]]
    for row in _rows(filtered_messages(f), lang, limit):
        data.append([Paragraph(_shape(str(row[i]))[:400], p) for i in keep])
    widths = [26 * mm, 24 * mm, 95 * mm, 18 * mm, 28 * mm, 26 * mm, 26 * mm, 25 * mm]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
    ]))
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def export_filename(ext: str) -> str:
    return f"messages_{local_today():%Y%m%d}_{uuidlib.uuid4().hex[:6]}.{ext}"
