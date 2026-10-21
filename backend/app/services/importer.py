"""Lecture et validation des fichiers Excel/CSV d'import (sans dépendance à la base).

Étapes :
1. read_table()      : lit le fichier (xlsx/csv), détecte la ligne d'en-tête et les colonnes.
2. suggest_mapping() : propose colonne -> variable (téléphone, message, nom, montant…).
3. validate_rows()   : normalise les numéros, détecte invalides / doublons / liste noire,
                       prépare les variables de fusion et le message personnalisé.
4. build_error_report() : rapport xlsx téléchargeable des lignes rejetées.
"""
import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .phone import error_label, normalize_phone
from .templating import normalize_key, render

PHONE_FIELD = "phone"
MESSAGE_FIELD = "message"
IGNORE_FIELD = ""

PHONE_HEADERS = {"telephone", "tel", "phone", "gsm", "mobile", "portable", "numero", "num", "numero_de_telephone",
                 "n_telephone", "numero_tel", "contact", "رقم", "الهاتف", "رقم_الهاتف", "الجوال"}
MESSAGE_HEADERS = {"message", "sms", "texte", "text", "contenu", "msg", "الرسالة", "رسالة", "النص"}

ROW_VALID = "valid"
ROW_INVALID = "invalid"
ROW_DUPLICATE = "duplicate"
ROW_BLACKLISTED = "blacklisted"


class ImportFileError(ValueError):
    """Fichier illisible ou vide."""


@dataclass
class Table:
    columns: list[str]
    rows: list[dict]          # {nom_colonne: valeur}
    sheet_name: str | None = None
    header_row: int = 1
    sheets: list[str] = field(default_factory=list)
    line_numbers: list[int] = field(default_factory=list)  # n° de ligne d'origine de chaque ligne de données

    def line_of(self, index: int) -> int:
        return self.line_numbers[index] if index < len(self.line_numbers) else self.header_row + index + 1


# --- Lecture -------------------------------------------------------------------

def read_table(source, filename: str, sheet_name: str | None = None, max_rows: int = 50000) -> Table:
    """source : chemin, bytes ou objet fichier."""
    ext = Path(filename).suffix.lower()
    data = _read_bytes(source)
    if not data:
        raise ImportFileError("empty_file")
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(data, sheet_name, max_rows)
    if ext in (".csv", ".txt"):
        return _read_csv(data, max_rows)
    raise ImportFileError("unsupported_format")


def _read_bytes(source) -> bytes:
    if isinstance(source, bytes):
        return source
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    return source.read()


def _clean_cell(value):
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _detect_header(raw_rows: list[list]) -> int:
    """Index (0-based) de la ligne d'en-tête parmi les 10 premières : la première dont au moins
    la moitié des cellules (par rapport à la ligne la plus large) sont du texte non numérique.
    Ignore ainsi les lignes de titre (« Export clients octobre ») placées au-dessus du tableau."""
    head = raw_rows[:10]
    width = max(sum(1 for c in r if c not in (None, "")) for r in head)
    for i, row in enumerate(head):
        texts = [c for c in row if isinstance(c, str) and c.strip() and not c.strip().lstrip("+").isdigit()]
        if texts and len(texts) * 2 >= width:
            return i
    return 0


def _build_table(raw_rows: list[list], max_rows: int, sheet_name=None, sheets=None) -> Table:
    numbered = [(i + 1, r) for i, r in enumerate(raw_rows) if any(c not in (None, "") for c in r)]
    if not numbered:
        raise ImportFileError("empty_file")
    raw_rows = [r for _, r in numbered]
    h = _detect_header(raw_rows)
    header = raw_rows[h]
    columns, seen = [], {}
    for idx, name in enumerate(header):
        name = str(name).strip() if name not in (None, "") else f"Colonne {idx + 1}"
        if name in seen:  # en-têtes en double : « Nom », « Nom (2) »
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        columns.append(name)
    body = numbered[h + 1:]
    if len(body) > max_rows:
        raise ImportFileError("too_many_rows")
    rows = []
    for _line, raw in body:
        rows.append({col: _clean_cell(raw[i]) if i < len(raw) else None for i, col in enumerate(columns)})
    return Table(columns, rows, sheet_name, numbered[h][0], sheets or [], [line for line, _ in body])


def _read_xlsx(data: bytes, sheet_name: str | None, max_rows: int) -> Table:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # fichier corrompu ou non xlsx
        raise ImportFileError("unreadable_file") from exc
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets[0]
        raw_rows = []
        for row in ws.iter_rows(values_only=True):
            raw_rows.append(list(row))
            if len(raw_rows) > max_rows + 20:
                raise ImportFileError("too_many_rows")
        return _build_table(raw_rows, max_rows, ws.title, wb.sheetnames)
    finally:
        wb.close()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1256", "cp1252"):  # cp1256 : CSV arabe exporté par Excel Windows
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_csv(data: bytes, max_rows: int) -> Table:
    text = _decode(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = []
    for row in reader:
        raw_rows.append(row)
        if len(raw_rows) > max_rows + 20:
            raise ImportFileError("too_many_rows")
    return _build_table(raw_rows, max_rows)


# --- Mappage -------------------------------------------------------------------

def suggest_mapping(table: Table, default_region: str = "MA") -> dict[str, str]:
    """Propose {colonne: champ}. Le téléphone est détecté par l'en-tête,
    sinon par le contenu (colonne avec le plus de numéros valides)."""
    mapping: dict[str, str] = {}
    phone_col = message_col = None
    for col in table.columns:
        key = normalize_key(col)
        if phone_col is None and (key in PHONE_HEADERS or key == "telephone"):
            phone_col = col
        elif message_col is None and key in MESSAGE_HEADERS:
            message_col = col
    if phone_col is None:
        sample = table.rows[:50]
        best, best_score = None, 0
        for col in table.columns:
            score = sum(1 for r in sample if normalize_phone(r.get(col), default_region).valid)
            if score > best_score:
                best, best_score = col, score
        if best and best_score >= max(1, len(sample) // 2):
            phone_col = best
    for col in table.columns:
        if col == phone_col:
            mapping[col] = PHONE_FIELD
        elif col == message_col:
            mapping[col] = MESSAGE_FIELD
        else:
            mapping[col] = normalize_key(col)
    return mapping


def phone_column(mapping: dict[str, str]) -> str | None:
    return next((c for c, f in mapping.items() if f == PHONE_FIELD), None)


# --- Validation ----------------------------------------------------------------

@dataclass
class ValidatedRow:
    row_number: int                    # numéro de ligne dans le fichier (1-based, en-tête inclus)
    raw: dict
    phone_raw: str | None
    phone_e164: str | None
    variables: dict
    custom_message: str | None
    status: str
    error: str | None = None
    rendered: str | None = None        # aperçu du message final

    @property
    def is_valid(self) -> bool:
        return self.status == ROW_VALID


@dataclass
class ValidationReport:
    rows: list[ValidatedRow]

    @property
    def total(self) -> int:
        return len(self.rows)

    def count(self, status: str) -> int:
        return sum(1 for r in self.rows if r.status == status)

    @property
    def valid(self) -> list[ValidatedRow]:
        return [r for r in self.rows if r.is_valid]

    @property
    def rejected(self) -> list[ValidatedRow]:
        return [r for r in self.rows if not r.is_valid]

    def summary(self) -> dict:
        return {"total": self.total, "valid": self.count(ROW_VALID), "invalid": self.count(ROW_INVALID),
                "duplicate": self.count(ROW_DUPLICATE), "blacklisted": self.count(ROW_BLACKLISTED)}


def _json_safe(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def validate_rows(table: Table, mapping: dict[str, str], body: str | None = None, *,
                  blacklist: set[str] | None = None, dedupe: bool = True,
                  default_region: str = "MA", lang: str = "fr") -> ValidationReport:
    """Valide chaque ligne. `body` est le texte commun (avec variables) ; une colonne
    mappée sur « message » le remplace ligne par ligne si elle est renseignée."""
    phone_col = phone_column(mapping)
    if phone_col is None:
        raise ValueError("phone_column_required")
    message_col = next((c for c, f in mapping.items() if f == MESSAGE_FIELD), None)
    blacklist = blacklist or set()
    seen: dict[str, int] = {}
    results = []

    for idx, row in enumerate(table.rows):
        row_number = table.line_of(idx)
        raw_json = {k: _json_safe(v) for k, v in row.items()}
        variables = {f: row.get(c) for c, f in mapping.items()
                     if f not in (PHONE_FIELD, MESSAGE_FIELD, IGNORE_FIELD)}
        custom = row.get(message_col) if message_col else None
        custom = str(custom).strip() if custom not in (None, "") else None
        phone = normalize_phone(row.get(phone_col), default_region)
        phone_raw = None if row.get(phone_col) is None else str(_json_safe(row.get(phone_col)))
        status, error, rendered = ROW_VALID, None, None

        if not phone.valid:
            status, error = ROW_INVALID, error_label(phone.error, lang)
        elif phone.e164 in blacklist:
            status, error = ROW_BLACKLISTED, ("Numéro en liste noire / désinscrit" if lang == "fr"
                                              else "رقم في القائمة السوداء")
        elif dedupe and phone.e164 in seen:
            status = ROW_DUPLICATE
            error = (f"Doublon de la ligne {seen[phone.e164]}" if lang == "fr"
                     else f"مكرر مع السطر {seen[phone.e164]}")
        template = custom or body
        if status == ROW_VALID:
            if not template:
                status, error = ROW_INVALID, "Message vide" if lang == "fr" else "رسالة فارغة"
            else:
                res = render(template, variables)
                rendered = res.text
                if not rendered.strip():
                    status, error = ROW_INVALID, "Message vide après fusion" if lang == "fr" else "رسالة فارغة"
                elif res.missing:
                    # Variable absente : on signale mais on n'empêche pas l'envoi si le texte reste non vide
                    error = ("Variables vides : " if lang == "fr" else "متغيرات فارغة: ") + ", ".join(res.missing)
        if phone.valid and phone.e164 not in seen:
            seen[phone.e164] = row_number

        results.append(ValidatedRow(row_number, raw_json, phone_raw, phone.e164,
                                    {k: _json_safe(v) for k, v in variables.items()},
                                    custom, status, error, rendered))
    return ValidationReport(results)


# --- Rapport d'erreurs -----------------------------------------------------------

def build_error_report(report: ValidationReport, columns: list[str], lang: str = "fr") -> bytes:
    labels = {
        "fr": {"line": "Ligne", "status": "Statut", "error": "Erreur", "summary": "Résumé",
               ROW_INVALID: "Invalide", ROW_DUPLICATE: "Doublon", ROW_BLACKLISTED: "Liste noire",
               ROW_VALID: "Valide", "total": "Total", "valid": "Valides", "invalid": "Invalides",
               "duplicate": "Doublons", "blacklisted": "Liste noire", "errors": "Erreurs"},
        "ar": {"line": "السطر", "status": "الحالة", "error": "الخطأ", "summary": "ملخص",
               ROW_INVALID: "غير صالح", ROW_DUPLICATE: "مكرر", ROW_BLACKLISTED: "قائمة سوداء",
               ROW_VALID: "صالح", "total": "المجموع", "valid": "صالحة", "invalid": "غير صالحة",
               "duplicate": "مكررة", "blacklisted": "قائمة سوداء", "errors": "الأخطاء"},
    }[lang if lang in ("fr", "ar") else "fr"]

    wb = Workbook()
    ws = wb.active
    ws.title = labels["errors"]
    if lang == "ar":
        ws.sheet_view.rightToLeft = True
    header = [labels["line"], labels["status"], labels["error"], *columns]
    ws.append(header)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="C0392B")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    for row in report.rejected:
        ws.append([row.row_number, labels.get(row.status, row.status), row.error,
                   *[row.raw.get(c) for c in columns]])
    ws.column_dimensions["C"].width = 40

    summary = wb.create_sheet(labels["summary"])
    if lang == "ar":
        summary.sheet_view.rightToLeft = True
    for key, value in report.summary().items():
        summary.append([labels[key], value])
    summary.column_dimensions["A"].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
